#!/usr/bin/env python3
"""Exercise the declared node-tree export action with synthetic content."""

import json
from io import BytesIO
import os
import runpy
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile
try:
    from docx import Document
except ImportError:
    print(json.dumps({"skipped": "needs python-docx"}))
    sys.exit(0)

ROOT = Path(__file__).resolve().parents[1]
RENDERER = ROOT / "do-work" / "render-structured-file-export.py"

WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": WORD_NS}


def node(text=None, children=None, heading_level=None):
    props = {"text": text} if text is not None else {}
    if heading_level is not None:
        props["heading_level"] = heading_level
    return {"props": props, "children": children or []}


def table_node(header, values):
    return node(children=[
        node(children=[node(value) for value in row])
        for row in (header, values)
    ])


def ordered_tree(title):
    return node(title, [node("Section heading", [
        node("Paragraph before the first table."),
        table_node(["First table header", "First table other"], ["First table value", "First table tail"]),
        node("Paragraph between the tables."),
        table_node(["Second table header", "Second table other"], ["Second table value", "Second table tail"]),
        node("Paragraph after the second table."),
        node("Subsection heading", [node("Paragraph in subsection.")], heading_level=2),
    ], heading_level=1)])


def assert_project_only_cli(directory):
    directory = Path(directory)
    destination = directory / "project-only-export-must-not-exist"
    result_path = directory / "project-only-result-must-not-exist.json"
    job = {
        "title": "",
        "descriptor": {
            "format_id": "pdf",
            "destination_dir": str(destination),
            "file_stem": "never-written",
        },
        "tree": node(children=[
            node("Main", [
                node("Opening paragraph."),
                node("Key points", [node("First item"), node("Second item")]),
                table_node(["Column A", "Column B"], ["Value A", "Value B"]),
                node("After table."),
            ], heading_level=1),
            node("Appendix", [node("Appendix detail.")], heading_level=2),
        ]),
    }
    rules = {
        "kind": "structured-file-projection-rules",
        "id": "project-only-test",
        "text_properties": ["text"],
        "table": {"min_cells": 2, "min_rows": 2},
        "skip_root_text_matching_title": True,
        "heading_level_path": ["props", "heading_level"],
        "formats": {"pdf": {
            "extension": "pdf",
            "renderer": "pdf",
            "font_family": "MissingForProjectOnly",
            "font_file": "not-installed-project-only-test.ttf",
            "default_heading_level": 2,
            "heading_styles": {
                "1": {"font_size": 18, "leading": 22, "space_before": 16, "space_after": 8},
                "2": {"font_size": 15, "leading": 19, "space_before": 14, "space_after": 7},
            },
        }},
    }
    job_path = directory / "project-only-job.json"
    rules_path = directory / "project-only-rules.json"
    job_path.write_text(json.dumps(job), encoding="utf-8")
    rules_path.write_text(json.dumps([{"payload": rules}]), encoding="utf-8")
    completed = subprocess.run([
        sys.executable, str(RENDERER),
        "--job-file", str(job_path),
        "--rules-rows-file", str(rules_path),
        "--result-path", str(result_path),
        "--project-only",
    ], cwd=ROOT, capture_output=True, text=True)
    assert completed.returncode == 0, completed.stderr or completed.stdout
    assert completed.stderr == ""
    projection = json.loads(completed.stdout)
    assert projection == {
        "kind": "structured-file-export-projection",
        "schema_version": 1,
        "format_id": "pdf",
        "title": "",
        "sections": [
            {
                "heading": "Main", "heading_level": 1,
                "lines": ["Opening paragraph."], "rows": [],
                "blocks": [{"type": "paragraph", "text": "Opening paragraph."}],
            },
            {
                "heading": "Key points", "heading_level": None,
                "lines": ["First item", "Second item"], "rows": [],
                "blocks": [
                    {"type": "paragraph", "text": "First item"},
                    {"type": "paragraph", "text": "Second item"},
                ],
            },
            {
                "heading": None, "heading_level": None,
                "lines": ["After table."],
                "rows": [["Column A", "Column B"], ["Value A", "Value B"]],
                "blocks": [
                    {"type": "table", "rows": [["Column A", "Column B"], ["Value A", "Value B"]]},
                    {"type": "paragraph", "text": "After table."},
                ],
            },
            {
                "heading": "Appendix", "heading_level": 2,
                "lines": ["Appendix detail."], "rows": [],
                "blocks": [{"type": "paragraph", "text": "Appendix detail."}],
            },
        ],
    }
    assert "payload_hash" not in projection
    assert not destination.exists()
    assert not result_path.exists()


def assert_paragraph_style_data(renderer):
    fmt = {"heading_styles": {"1": {"style_id": "Heading1", "name": "heading 1", "size_half_points": 30,
                                    "space_before_twips": 240, "space_after_twips": 100}},
           "body_style": {"font_family": "Declared Script Family", "size_half_points": 22,
                          "space_after_twips": 120, "line_spacing": 276, "widow_control": True},
           "title_style": {"space_after_twips": 180, "widow_control": False},
           "table_style": {"size_half_points": 20, "space_after_twips": 0, "line_spacing": 240},
           "table_header_rows": 1, "space_after_table_twips": 120, "table_allow_row_split": False}
    sections = [{"heading": "Section", "heading_level": 1, "blocks": [
        {"type": "paragraph", "text": "First paragraph."},
        {"type": "paragraph", "text": "Second paragraph."},
        {"type": "table", "rows": [["Key", "Value"], ["One", "Two"]]},
        {"type": "paragraph", "text": "After the table."},
    ]}]
    document = Document(BytesIO(renderer["render_docx"]("Title", sections, fmt)))
    normal = document.styles["Normal"]
    assert normal.font.name == "Declared Script Family"
    fonts = normal._element.rPr.rFonts
    assert all(fonts.get(f"{{{WORD_NS}}}{script}") == "Declared Script Family"
               for script in ("ascii", "hAnsi", "eastAsia", "cs"))
    assert normal.font.size.pt == 11
    assert normal.paragraph_format.space_after.pt == 6
    assert normal.paragraph_format.line_spacing == 1.15
    assert normal.paragraph_format.widow_control is True
    assert document.styles["heading 1"].paragraph_format.space_before.pt == 12
    assert document.styles["heading 1"].paragraph_format.space_after.pt == 5
    assert document.styles["Title"].paragraph_format.space_after.pt == 9
    assert document.styles["Title"].paragraph_format.widow_control is False
    unconfigured = Document(BytesIO(renderer["render_docx"]("", [], {})))
    assert unconfigured.styles["Normal"].paragraph_format.widow_control is None
    cell = document.tables[0].cell(0, 0).paragraphs[0]
    assert cell.style.style_id == "TableText"
    assert cell.style.font.size.pt == 10
    assert cell.style.paragraph_format.space_after.pt == 0
    assert [p.text for p in document.paragraphs] == ["Title", "Section", "First paragraph.", "Second paragraph.", "After the table."]
    assert document.paragraphs[-1].paragraph_format.space_before.pt == 6
    assert document.paragraphs[-2].paragraph_format.space_before is None
    assert document.tables[0].rows[0]._tr.find("w:trPr/w:tblHeader", NS) is not None
    assert document.tables[0].rows[1]._tr.find("w:trPr/w:tblHeader", NS) is None
    assert all(row._tr.find("w:trPr/w:cantSplit", NS) is not None for row in document.tables[0].rows)
    for invalid in ({"body_style": {"space_after_twips": True}},
                    {"body_style": {"line_spacing": 0}},
                    {"body_style": {"widow_control": 1}},
                    {"body_style": {"widow_control": "true"}},
                    {"title_style": {"font_family": ""}},
                    {"heading_styles": {"1": {"style_id": "Normal", "size_half_points": 30}}},
                    {"table_header_rows": True}, {"table_header_rows": 3},
                    {"space_after_table_twips": -1}, {"table_allow_row_split": "no"}):
        try:
            renderer["render_docx"]("Title", sections, {**fmt, **invalid})
        except ValueError:
            pass
        else:
            raise AssertionError("Invalid style DATA was accepted")


def assert_no_break_character_data(renderer):
    title = "Periods 3–8"
    paragraph = "Use 12–24 units, 7–10 days and 15–30 minutes. Keep A—B and 1-2 unchanged."
    sections = [{"heading": "Chapter 2–4", "heading_level": 1, "blocks": [
        {"type": "paragraph", "text": paragraph},
        {"type": "table", "rows": [["Slot", "Range"], ["A/B", "40–60"]]},
    ]}]
    original = json.dumps(sections, ensure_ascii=False)
    fmt = {"heading_styles": {"1": {"style_id": "Heading1", "size_half_points": 30}}}
    baseline = Document(BytesIO(renderer["render_docx"](title, sections, fmt)))
    protected = Document(BytesIO(renderer["render_docx"](
        title, sections, {**fmt, "no_break_characters": ["–", "/"]})))
    def texts(document):
        return ([p.text for p in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    assert texts(baseline) == [title, "Chapter 2–4", paragraph, "Slot", "Range", "A/B", "40–60"]
    assert [text.replace("\u2060", "") for text in texts(protected)] == texts(baseline)
    assert all("\u2060–\u2060" in text for text in texts(protected) if "–" in text)
    assert protected.tables[0].cell(1, 0).text == "A\u2060/\u2060B"
    assert "A—B and 1-2 unchanged" in protected.paragraphs[2].text
    assert json.dumps(sections, ensure_ascii=False) == original
    # Existing formatting controls and adjacent configured characters remain
    # idempotent; XML-sensitive configured punctuation remains escaped.
    for value, characters in [("x\u2060–\u2060y", ["–"]), ("a//b", ["/"]), ("a&b<c", ["&", "<"])]:
        first = Document(BytesIO(renderer["render_docx"](value, [], {"no_break_characters": characters})))
        text = first.paragraphs[0].text
        again = Document(BytesIO(renderer["render_docx"](text, [], {"no_break_characters": characters})))
        assert again.paragraphs[0].text == text
        assert text.replace("\u2060", "") == value.replace("\u2060", "")
    for invalid in (None, "–", [True], [""], ["12–24"], [" "], ["\n"], ["\u2060"], ["x"] * 65):
        try:
            renderer["render_docx"]("", [], {"no_break_characters": invalid})
        except ValueError:
            pass
        else:
            raise AssertionError("Invalid no-break DATA was accepted")


def main():
    assert_no_break_character_data(runpy.run_path(str(RENDERER)))
    with tempfile.TemporaryDirectory(prefix="structured-output-docx-") as directory:
        assert_project_only_cli(directory)
        title = "Synthetic Word export"
        job = {
            "title": title,
            "descriptor": {"format_id": "docx", "destination_dir": directory, "file_stem": "synthetic"},
            "tree": ordered_tree(title),
        }
        rules = {
            "kind": "structured-file-projection-rules",
            "id": "synthetic",
            "text_properties": ["text"],
            "table": {"min_cells": 2, "min_rows": 2},
            "skip_root_text_matching_title": True,
            "heading_level_path": ["props", "heading_level"],
            "formats": {"docx": {
                "extension": "docx", "renderer": "docx", "default_heading_level": 1,
                "heading_styles": {
                    "1": {"style_id": "Heading1", "size_half_points": 32},
                    "2": {"style_id": "Heading2", "size_half_points": 28},
                },
            }},
        }
        renderer = runpy.run_path(str(RENDERER))
        assert_paragraph_style_data(renderer)
        sections = renderer["project"](job["tree"], rules, title=title)
        assert sections[0]["lines"] == [
            "Paragraph before the first table.", "Paragraph between the tables.",
            "Paragraph after the second table.",
        ]
        assert sections[0]["rows"] == [
            ["First table header", "First table other"],
            ["First table value", "First table tail"],
            ["Second table header", "Second table other"],
            ["Second table value", "Second table tail"],
        ]
        assert [block["type"] for block in sections[0]["blocks"]] == [
            "paragraph", "table", "paragraph", "table", "paragraph",
        ]
        assert [section["heading_level"] for section in sections] == [1, 2]
        result_path = Path(directory) / "result.json"
        job_path = Path(directory) / "job.json"
        rules_path = Path(directory) / "rules-rows.json"
        job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
        rules_path.write_text(json.dumps([{"payload": rules}], ensure_ascii=False), encoding="utf-8")
        subprocess.run([
            sys.executable,
            str(RENDERER),
            "--job-file", str(job_path),
            "--rules-rows-file", str(rules_path),
            "--result-path", str(result_path),
        ], check=True, cwd=ROOT, capture_output=True, text=True)
        result = json.loads(result_path.read_text())
        path = Path(result["path"])
        assert path.suffix == ".docx"
        assert result["format_id"] == "docx"
        with ZipFile(path) as docx:
            assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml", "word/styles.xml"}.issubset(docx.namelist())
            document = ET.fromstring(docx.read("word/document.xml"))
        body = document.find("w:body", NS)
        paragraphs = body.findall("w:p", NS)
        sequence = []
        for element in body:
            if element.tag == f"{{{WORD_NS}}}p":
                sequence.append("".join(element.itertext()))
            elif element.tag == f"{{{WORD_NS}}}tbl":
                sequence.append(["".join(cell.itertext()) for cell in element.findall("w:tr/w:tc", NS)])
        assert sequence == [
            title, "Section heading", "Paragraph before the first table.",
            ["First table header", "First table other", "First table value", "First table tail"],
            "Paragraph between the tables.",
            ["Second table header", "Second table other", "Second table value", "Second table tail"],
            "Paragraph after the second table.", "Subsection heading", "Paragraph in subsection.",
        ]
        rows = document.findall(".//w:body/w:tbl/w:tr", NS)
        assert len(rows) == 4
        heading = paragraphs[1]
        assert heading.find("w:pPr/w:keepNext", NS) is not None
        assert heading.find("w:pPr/w:keepLines", NS) is not None
        with ZipFile(path) as docx:
            styles = ET.fromstring(docx.read("word/styles.xml"))
        heading_style = styles.find(".//w:style[@w:styleId='Heading1']", NS)
        assert heading_style is not None
        assert heading_style.find("w:pPr/w:outlineLvl", NS).attrib[f"{{{WORD_NS}}}val"] == "0"
        subsection_style = styles.find(".//w:style[@w:styleId='Heading2']", NS)
        assert subsection_style is not None
        assert subsection_style.find("w:pPr/w:outlineLvl", NS).attrib[f"{{{WORD_NS}}}val"] == "1"
        assert [
            p.find("w:pPr/w:pStyle", NS).attrib[f"{{{WORD_NS}}}val"]
            for p in paragraphs if p.find("w:pPr/w:pStyle", NS) is not None
        ] == ["Title", "Heading1", "Heading2"]
        consumer = Document(path)
        assert [paragraph.style.style_id for paragraph in consumer.paragraphs] == [
            "Title", "Heading1", "Normal", "Normal", "Normal", "Heading2", "Normal",
        ]
        assert [paragraph.text for paragraph in consumer.paragraphs if paragraph.style.style_id.startswith("Heading")] == [
            "Section heading", "Subsection heading",
        ]
        readback = Path(__file__).with_name("read-structured-docx.py")
        inspection = json.loads(subprocess.run(
            [sys.executable, str(readback), str(path)], check=True, cwd=ROOT,
            capture_output=True, text=True,
        ).stdout)
        assert inspection["counts"] == {
            "paragraphs": 7, "headings": 2, "lists": 0, "tables": 2,
            "table_rows": 4, "cells": 8,
        }
        assert [block.get("heading_level") for block in inspection["blocks"] if "heading_level" in block] == [1, 2]
        assert [block["kind"] for block in inspection["blocks"]] == [
            "paragraph", "paragraph", "paragraph", "table", "paragraph", "table",
            "paragraph", "paragraph", "paragraph",
        ]

        # The plain text renderer must traverse the same ordered projection.
        text_job = {
            "title": title,
            "descriptor": {"format_id": "txt", "destination_dir": directory, "file_stem": "synthetic"},
            "tree": ordered_tree(title),
        }
        text_rules = {
            **rules,
            "formats": {"txt": {
                "extension": "txt", "renderer": "text", "title_template": "{title}",
                "heading_template": "{heading}", "line_template": "{line}",
                "row_template": "{cells}", "part_separator": "\n\n",
                "document_template": "{body}\n",
            }},
        }
        text_result = Path(directory) / "text-result.json"
        subprocess.run([
            sys.executable, str(RENDERER),
            "--job-json", json.dumps(text_job),
            "--rules-rows-json", json.dumps([{"payload": text_rules}]),
            "--result-path", str(text_result),
        ], check=True, cwd=ROOT, capture_output=True, text=True)
        text_path = Path(json.loads(text_result.read_text())["path"])
        text = text_path.read_text()
        ordered_text = [
            "Section heading", "Paragraph before the first table.", "First table header",
            "First table value", "Paragraph between the tables.", "Second table header",
            "Second table value", "Paragraph after the second table.",
        ]
        positions = [text.index(value) for value in ordered_text]
        assert positions == sorted(positions), text

        # File inputs support a real job larger than this host's entire argv
        # budget; no manuscript or other caller content is used here.
        arg_max = os.sysconf("SC_ARG_MAX") if hasattr(os, "sysconf") else 1_048_576
        large_text = "x" * max(1_100_000, arg_max + 4096)
        large_job = {
            "title": "Large synthetic file input",
            "descriptor": {"format_id": "docx", "destination_dir": directory, "file_stem": "large-synthetic"},
            "tree": node("Large synthetic file input", [node("Large section", [node(large_text)], heading_level=1)]),
        }
        large_job_path = Path(directory) / "large-job.json"
        large_job_path.write_text(json.dumps(large_job), encoding="utf-8")
        assert large_job_path.stat().st_size > arg_max
        assert large_job_path.stat().st_size < renderer["MAX_JSON_INPUT_BYTES"]
        large_result_path = Path(directory) / "large-result.json"
        subprocess.run([
            sys.executable, str(RENDERER),
            "--job-file", str(large_job_path),
            "--rules-rows-file", str(rules_path),
            "--result-path", str(large_result_path),
        ], check=True, cwd=ROOT, capture_output=True, text=True)
        large_result = json.loads(large_result_path.read_text())
        large_docx_path = Path(large_result["path"])
        with ZipFile(large_docx_path) as large_docx:
            large_xml = ET.fromstring(large_docx.read("word/document.xml"))
        assert any(len(element.text or "") == len(large_text) for element in large_xml.findall(".//w:t", NS))

        # Empty title suppresses only the optional title paragraph; an H1 at
        # the root of the projected source remains the document's first block.
        empty_title_job = {
            "title": "",
            "descriptor": {"format_id": "docx", "destination_dir": directory, "file_stem": "empty-title"},
            "tree": node(children=[node("Manuscript", [node("Opening paragraph.")], heading_level=1)]),
        }
        empty_title_path = Path(directory) / "empty-title-job.json"
        empty_title_path.write_text(json.dumps(empty_title_job), encoding="utf-8")
        empty_title_result_path = Path(directory) / "empty-title-result.json"
        subprocess.run([
            sys.executable, str(RENDERER),
            "--job-file", str(empty_title_path),
            "--rules-rows-file", str(rules_path),
            "--result-path", str(empty_title_result_path),
        ], check=True, cwd=ROOT, capture_output=True, text=True)
        empty_title_result = json.loads(empty_title_result_path.read_text())
        with ZipFile(empty_title_result["path"]) as empty_title_docx:
            empty_title_xml = ET.fromstring(empty_title_docx.read("word/document.xml"))
        empty_title_paragraphs = empty_title_xml.findall(".//w:body/w:p", NS)
        empty_title_styles = [paragraph.find("w:pPr/w:pStyle", NS) for paragraph in empty_title_paragraphs]
        assert empty_title_styles[0].attrib[f"{{{WORD_NS}}}val"] == "Heading1"
        assert ["".join(paragraph.itertext()) for paragraph in empty_title_paragraphs] == [
            "Manuscript", "Opening paragraph.",
        ]

        def fail_safely(command, expected):
            failed = subprocess.run(command, cwd=ROOT, capture_output=True, text=True)
            output = failed.stdout + failed.stderr
            assert failed.returncode != 0, output
            assert expected in output, output
            assert "PRIVATE_INPUT_MARKER" not in output, output

        fail_safely([
            sys.executable, str(RENDERER), "--job-file", str(job_path),
            "--job-json", '{"text":"PRIVATE_INPUT_MARKER"}',
            "--rules-rows-file", str(rules_path),
        ], "not allowed with argument")
        fail_safely([
            sys.executable, str(RENDERER), "--job-file", str(job_path),
            "--rules-rows-file", str(rules_path), "--rules-rows-json", "[]",
        ], "not allowed with argument")

        malformed_job = Path(directory) / "malformed-job.json"
        malformed_job.write_text('{"text":"PRIVATE_INPUT_MARKER"', encoding="utf-8")
        fail_safely([
            sys.executable, str(RENDERER), "--job-file", str(malformed_job),
            "--rules-rows-file", str(rules_path),
        ], "job file is not valid JSON")

        invalid_utf8_job = Path(directory) / "invalid-utf8-job.json"
        invalid_utf8_job.write_bytes(b'{"text":"' + b"\xff" + b'"}')
        fail_safely([
            sys.executable, str(RENDERER), "--job-file", str(invalid_utf8_job),
            "--rules-rows-file", str(rules_path),
        ], "job file is not valid UTF-8")

        malformed_rules = Path(directory) / "malformed-rules.json"
        malformed_rules.write_text('[{"payload":"PRIVATE_INPUT_MARKER"', encoding="utf-8")
        fail_safely([
            sys.executable, str(RENDERER), "--job-file", str(job_path),
            "--rules-rows-file", str(malformed_rules),
        ], "projection rules file is not valid JSON")

        too_large_job = Path(directory) / "oversized-job.json"
        too_large_job.write_bytes(b" " * (renderer["MAX_JSON_INPUT_BYTES"] + 1))
        fail_safely([
            sys.executable, str(RENDERER), "--job-file", str(too_large_job),
            "--rules-rows-file", str(rules_path),
        ], "exceeds the")

        linked_job = Path(directory) / "linked-job.json"
        os.symlink(job_path, linked_job)
        fail_safely([
            sys.executable, str(RENDERER), "--job-file", str(linked_job),
            "--rules-rows-file", str(rules_path),
        ], "regular, non-symlink file")

        non_file_job = Path(directory) / "job-directory"
        non_file_job.mkdir()
        fail_safely([
            sys.executable, str(RENDERER), "--job-file", str(non_file_job),
            "--rules-rows-file", str(rules_path),
        ], "regular, non-symlink file")

        def reject_during_read(path, change_file):
            real_fstat = os.fstat
            calls = 0

            def change_after_open(fd):
                nonlocal calls
                snapshot = real_fstat(fd)
                if calls == 0:
                    calls += 1
                    change_file()
                return snapshot

            os.fstat = change_after_open
            try:
                try:
                    renderer["read_json_file"](str(path), "job")
                except ValueError as exc:
                    message = str(exc)
                    assert message == "job file changed during read or was replaced", message
                    assert "PRIVATE_INPUT_MARKER" not in message
                else:
                    raise AssertionError("reader accepted a file changed during read")
            finally:
                os.fstat = real_fstat

        changing_job = Path(directory) / "changing-job.json"
        changing_job.write_text('{"value":"before"}', encoding="utf-8")
        reject_during_read(
            changing_job,
            lambda: changing_job.write_text('{"value":"PRIVATE_INPUT_MARKER changed after open"}', encoding="utf-8"),
        )

        replaced_job = Path(directory) / "replaced-job.json"
        replacement_bytes = Path(directory) / "replacement-bytes.json"
        replaced_job.write_text('{"value":"before"}', encoding="utf-8")
        replacement_bytes.write_text('{"value":"replacement"}', encoding="utf-8")
        reject_during_read(replaced_job, lambda: os.replace(replacement_bytes, replaced_job))


if __name__ == "__main__":
    main()
